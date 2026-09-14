"""Penulis berkas DOCX ekspor bab."""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt

from app.chapter_export.renderers.base import iter_export_chapters, publish_rendering_stage


async def render_docx(context, payload: dict[str, Any], part_path: Path) -> None:
    """Merakit dokumen Word lalu menyimpannya di utas terpisah.

    Berbeda dari TXT, python-docx menyusun seluruh dokumen di memori sebelum dapat disimpan,
    sehingga penyimpanan dijalankan lewat utas agar perulangan peristiwa tidak terhalang.
    Judul memakai gaya Heading bawaan Word supaya panel navigasi dokumen langsung berfungsi.
    """
    document = Document()
    body_style = document.styles["Normal"]
    font = getattr(body_style, "font", None)
    if font is not None:
        font.name = "Calibri"
        font.size = Pt(11)

    total = 0
    async for chapter in iter_export_chapters(context, payload):
        if chapter.volume_heading is not None:
            heading = document.add_heading(chapter.volume_heading, level=1)
            if not chapter.is_first:
                heading.paragraph_format.page_break_before = True
        chapter_heading = document.add_heading(chapter.title, level=2)
        if chapter.volume_heading is None and not chapter.is_first:
            chapter_heading.paragraph_format.page_break_before = True
        # Isi bab tersimpan sebagai teks polos berbasis baris, jadi setiap baris menjadi satu
        # paragraf agar jeda yang ditulis pengarang tetap terlihat di Word.
        for line in chapter.content.split("\n"):
            document.add_paragraph(line)
        total += 1

    await context.check_cancelled()
    await publish_rendering_stage(context, total)
    await asyncio.to_thread(document.save, str(part_path))
