# -*- coding: utf-8 -*-
"""Pengujian API ekspor bab."""

from datetime import UTC, datetime, timedelta
import json

import pytest
from docx import Document
from httpx import AsyncClient
from urllib.parse import unquote

from app.background.events.publisher import BackgroundEventPublisher
from app.background.jobs import service as background_service
from app.background.runtime.context import JobContext
from app.background.runtime.dispatcher import dispatch_job
from app.chapter_export import service as chapter_export_service
from app.chapter_export.renderers import pdf as pdf_renderer
from app.background.jobs.models import BackgroundJob
from app.api.routers import chapter_exports as chapter_exports_router
from app.storage.repos import chapter_repo


def test_volume_numbers() -> None:
    assert chapter_export_service.volume_number(1) == "1"
    assert chapter_export_service.volume_number(10) == "10"
    assert chapter_export_service.volume_number(11) == "11"
    assert chapter_export_service.volume_number(21) == "21"
    assert chapter_export_service.volume_number(101) == "101"


def test_volume_export_heading_keeps_stored_title() -> None:
    """Judul volume sudah memuat penomorannya sendiri, jadi ekspor tidak menambah awalan lagi."""
    assert chapter_export_service.volume_export_heading(1, "Volume 1") == "Volume 1"
    assert chapter_export_service.volume_export_heading(2, "Volume Satu") == "Volume Satu"
    assert chapter_export_service.volume_export_heading(3, "Bagian Pendahuluan") == (
        "Bagian Pendahuluan"
    )


def test_volume_export_heading_falls_back_to_order_when_title_blank() -> None:
    """Volume tanpa judul tetap memerlukan penanda agar batas antarvolume tidak hilang."""
    assert chapter_export_service.volume_export_heading(4, "") == "Volume 4"
    assert chapter_export_service.volume_export_heading(5, "   ") == "Volume 5"


def test_expired_export_is_not_downloadable(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job = BackgroundJob(
        id="expired-export",
        type=chapter_export_service.EXPORT_JOB_TYPE,
        status="succeeded",
        payload_json='{"filename":"uji.txt"}',
        result_json='{"expires_at":"2020-01-01T00:00:00+00:00"}',
    )
    _part_path, output_path = chapter_export_service.export_file_paths(job.id)
    output_path.write_text("expired", encoding="utf-8")

    assert not chapter_export_service.is_export_download_available(job)


async def _create_project(
    client: AsyncClient,
    title: str = "Novel Uji",
    book_type: str = "fiction",
) -> tuple[str, str]:
    response = await client.post(
        "/api/v1/projects",
        data={"title": title, "book_type": book_type},
    )
    assert response.status_code == 201
    project_id = response.json()["id"]
    volumes = (await client.get(f"/api/v1/projects/{project_id}/volumes")).json()
    return project_id, volumes[0]["id"]


async def _create_chapter(
    client: AsyncClient,
    project_id: str,
    volume_id: str,
    title: str,
    content: str,
    word_count: int,
) -> dict:
    response = await client.post(
        f"/api/v1/projects/{project_id}/chapters",
        json={
            "volume_id": volume_id,
            "title": title,
            "content": content,
            "word_count": word_count,
        },
    )
    assert response.status_code == 201
    return response.json()


@pytest.mark.asyncio
async def test_create_full_volume_export_uses_volume_filename_and_snapshot_selection(
    client: AsyncClient,
) -> None:
    project_id, volume_id = await _create_project(client)
    first = await _create_chapter(client, project_id, volume_id, "Bab 1", "Isi bab 1\r\nBaris kedua", 5)
    second = await _create_chapter(client, project_id, volume_id, "Bab 2", "Isi bab 2", 5)

    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 201
    data = response.json()
    assert data["status"] == "pending"
    assert data["chapter_count"] == 2
    assert data["word_count"] == 10
    assert data["filename"] == "Novel Uji-Lengkap-2026-07-28.txt"
    assert data["chapter_ids"] == [first["id"], second["id"]]


@pytest.mark.asyncio
async def test_export_creation_does_not_load_chapter_bodies(client: AsyncClient, monkeypatch) -> None:
    project_id, volume_id = await _create_project(client)
    chapter = await _create_chapter(client, project_id, volume_id, "Bab 1", "Isi utama", 2)

    async def reject_full_chapter_load(*_args, **_kwargs):
        raise AssertionError("Tahap pembuatan ekspor tidak boleh membaca isi utama bab secara lengkap")

    monkeypatch.setattr(chapter_repo, "list_by_project", reject_full_chapter_load)
    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [chapter["id"]],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 201


@pytest.mark.asyncio
async def test_only_cancel_endpoint_preempts_running_export(client: AsyncClient, monkeypatch) -> None:
    class Supervisor:
        def __init__(self) -> None:
            self.cancelled_job_ids: list[str] = []

        def create_event_publisher(self) -> BackgroundEventPublisher:
            return BackgroundEventPublisher(None)

        def cancel_running_chapter_export(self, job_id: str) -> bool:
            self.cancelled_job_ids.append(job_id)
            return False

    supervisor = Supervisor()
    monkeypatch.setattr(chapter_exports_router, "get_background_supervisor", lambda: supervisor)
    project_id, volume_id = await _create_project(client)
    chapter = await _create_chapter(client, project_id, volume_id, "Bab 1", "Isi utama", 2)

    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [chapter["id"]],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )
    assert created.status_code == 201
    assert supervisor.cancelled_job_ids == []

    cancelled = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports/{created.json()['id']}/cancel"
    )
    assert cancelled.status_code == 200
    assert supervisor.cancelled_job_ids == [created.json()["id"]]


@pytest.mark.asyncio
async def test_create_fragment_export_uses_chapter_filename(client: AsyncClient) -> None:
    project_id, volume_id = await _create_project(client)
    selected = await _create_chapter(client, project_id, volume_id, "Bab 1", "Isi utama", 2)
    await _create_chapter(client, project_id, volume_id, "Bab 2", "Isi utama", 2)

    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [selected["id"]],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 201
    assert response.json()["filename"] == "Novel Uji-1 Bab-2026-07-28.txt"


@pytest.mark.asyncio
async def test_manually_selected_complete_volume_still_uses_chapter_format(
    client: AsyncClient,
) -> None:
    project_id, volume_id = await _create_project(client)
    first = await _create_chapter(client, project_id, volume_id, "Bab 1", "Isi utama", 2)
    second = await _create_chapter(client, project_id, volume_id, "Bab 2", "Isi utama", 2)

    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [first["id"], second["id"]],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 201
    assert response.json()["filename"] == "Novel Uji-2 Bab-2026-07-28.txt"


@pytest.mark.asyncio
async def test_create_export_rejects_empty_selection(client: AsyncClient) -> None:
    project_id, _volume_id = await _create_project(client)

    response = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )

    assert response.status_code == 400
    assert "bab" in response.json()["detail"]


@pytest.mark.asyncio
async def test_export_task_writes_full_volume_txt_and_serves_download(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(client)
    first = await _create_chapter(client, project_id, volume_id, "Bab 1", "Isi bab 1\r\nBaris kedua", 5)
    second = await _create_chapter(client, project_id, volume_id, "Bab 2", "Isi bab 2", 5)
    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )
    assert created.status_code == 201
    job = await background_service.get_job(session, created.json()["id"])
    assert job is not None
    job.status = "running"
    await session.commit()

    context = JobContext(session=session, job=job, publisher=BackgroundEventPublisher(None))
    result = await dispatch_job(context)
    await background_service.mark_succeeded(session, context.publisher, context.job, result=result)
    await session.commit()

    status_response = await client.get(
        f"/api/v1/projects/{project_id}/chapter-exports/{job.id}"
    )
    assert status_response.status_code == 200
    assert status_response.json()["current"] == 2
    assert status_response.json()["total"] == 2
    assert status_response.json()["download_url"]

    download_response = await client.get(
        f"/api/v1/projects/{project_id}/chapter-exports/{job.id}/download"
    )
    assert download_response.status_code == 200
    assert "attachment" in download_response.headers["content-disposition"]
    assert "Novel Uji-Lengkap-2026-07-28.txt" in unquote(
        download_response.headers["content-disposition"]
    )
    assert download_response.content.decode("utf-8-sig") == (
        "Volume 1\nBab 1\nIsi bab 1\nBaris kedua\n\nBab 2\nIsi bab 2"
    )
    assert result == {
        "filename": "Novel Uji-Lengkap-2026-07-28.txt",
        "format": "txt",
        "volume_count": 1,
        "chapter_count": 2,
        "word_count": 10,
        "expires_at": result["expires_at"],
    }
    assert [first["id"], second["id"]] == created.json()["chapter_ids"]


@pytest.mark.asyncio
async def test_cancelled_export_removes_partial_file(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    project_id, volume_id = await _create_project(client)
    await _create_chapter(client, project_id, volume_id, "Bab 1", "Isi utama", 2)
    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
        },
    )
    job_id = created.json()["id"]
    part_path, output_path = chapter_export_service.export_file_paths(job_id)
    part_path.write_text("partial", encoding="utf-8")
    output_path.write_text("complete", encoding="utf-8")

    cancelled = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports/{job_id}/cancel"
    )

    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled"
    assert not part_path.exists()
    assert not output_path.exists()


@pytest.mark.asyncio
async def test_cleanup_keeps_output_while_export_is_still_running(
    session,
    monkeypatch,
    tmp_path,
) -> None:
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    job = BackgroundJob(
        id="running-export",
        type=chapter_export_service.EXPORT_JOB_TYPE,
        status="running",
        payload_json="{}",
    )
    session.add(job)
    await session.commit()
    _part_path, output_path = chapter_export_service.export_file_paths(job.id)
    output_path.write_text("finished but not committed", encoding="utf-8")

    assert await chapter_export_service.cleanup_chapter_export_files(session) == 0
    assert output_path.exists()


@pytest.mark.asyncio
async def test_cleanup_keeps_succeeded_output_until_ttl_expires(
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Berkas hasil tugas sukses harus bertahan selama masa berlaku unduhan belum lewat."""
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    expires_at = datetime.now(UTC) + timedelta(hours=12)
    job = BackgroundJob(
        id="fresh-succeeded-export",
        type=chapter_export_service.EXPORT_JOB_TYPE,
        status="succeeded",
        payload_json='{"filename":"uji.txt"}',
        result_json=json.dumps({"expires_at": expires_at.isoformat()}),
    )
    session.add(job)
    await session.commit()
    _part_path, output_path = chapter_export_service.export_file_paths(job.id)
    output_path.write_text("hasil ekspor", encoding="utf-8")

    assert await chapter_export_service.cleanup_chapter_export_files(session) == 0
    assert output_path.exists()
    assert chapter_export_service.is_export_download_available(job)


@pytest.mark.asyncio
async def test_cleanup_removes_succeeded_output_after_ttl_expires(
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Berkas hasil tugas sukses harus disapu setelah masa berlaku unduhan lewat."""
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    expires_at = datetime.now(UTC) - timedelta(minutes=1)
    job = BackgroundJob(
        id="stale-succeeded-export",
        type=chapter_export_service.EXPORT_JOB_TYPE,
        status="succeeded",
        payload_json='{"filename":"uji.txt"}',
        result_json=json.dumps({"expires_at": expires_at.isoformat()}),
    )
    session.add(job)
    await session.commit()
    _part_path, output_path = chapter_export_service.export_file_paths(job.id)
    output_path.write_text("hasil kedaluwarsa", encoding="utf-8")

    assert await chapter_export_service.cleanup_chapter_export_files(session) == 1
    assert not output_path.exists()


@pytest.mark.asyncio
async def test_cleanup_removes_leftover_part_file_of_succeeded_export(
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Berkas sementara milik tugas sukses adalah sisa gagal rename dan harus dihapus."""
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    expires_at = datetime.now(UTC) + timedelta(hours=12)
    job = BackgroundJob(
        id="succeeded-with-part",
        type=chapter_export_service.EXPORT_JOB_TYPE,
        status="succeeded",
        payload_json='{"filename":"uji.txt"}',
        result_json=json.dumps({"expires_at": expires_at.isoformat()}),
    )
    session.add(job)
    await session.commit()
    part_path, output_path = chapter_export_service.export_file_paths(job.id)
    part_path.write_text("sisa berkas sementara", encoding="utf-8")
    output_path.write_text("hasil ekspor", encoding="utf-8")

    assert await chapter_export_service.cleanup_chapter_export_files(session) == 1
    assert not part_path.exists()
    assert output_path.exists()


async def _run_export_job(
    client: AsyncClient,
    session,
    project_id: str,
    volume_id: str,
    export_format: str,
) -> tuple[dict, BackgroundJob]:
    """Menjalankan satu tugas ekspor sampai selesai lalu mengembalikan hasil dan catatan tugasnya."""
    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
            "format": export_format,
        },
    )
    assert created.status_code == 201
    assert created.json()["format"] == export_format
    job = await background_service.get_job(session, created.json()["id"])
    assert job is not None
    job.status = "running"
    await session.commit()

    context = JobContext(session=session, job=job, publisher=BackgroundEventPublisher(None))
    result = await dispatch_job(context)
    assert result is not None
    await background_service.mark_succeeded(session, context.publisher, context.job, result=result)
    await session.commit()
    return result, job


@pytest.mark.asyncio
async def test_export_task_writes_docx_and_serves_download(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(client)
    await _create_chapter(client, project_id, volume_id, "Bab 1", "Isi bab 1\r\nBaris kedua", 5)
    await _create_chapter(client, project_id, volume_id, "Bab 2", "Isi bab 2", 5)

    result, job = await _run_export_job(client, session, project_id, volume_id, "docx")

    assert result["filename"] == "Novel Uji-Lengkap-2026-07-28.docx"
    assert result["format"] == "docx"
    _part_path, output_path = chapter_export_service.export_file_paths(job.id, "docx")
    assert output_path.exists()

    download_response = await client.get(
        f"/api/v1/projects/{project_id}/chapter-exports/{job.id}/download"
    )
    assert download_response.status_code == 200
    assert download_response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert download_response.content[:4] == b"PK\x03\x04"
    assert "Novel Uji-Lengkap-2026-07-28.docx" in unquote(
        download_response.headers["content-disposition"]
    )

    document = Document(str(output_path))
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert "Volume 1" in texts
    assert texts.index("Bab 1") < texts.index("Bab 2")
    assert "Isi bab 1" in texts
    assert "Baris kedua" in texts
    assert "Isi bab 2" in texts
    styles = {
        paragraph.text: getattr(paragraph.style, "name", None)
        for paragraph in document.paragraphs
        if paragraph.text
    }
    assert styles["Volume 1"] == "Heading 1"
    assert styles["Bab 1"] == "Heading 2"


@pytest.mark.asyncio
async def test_export_task_writes_pdf_and_serves_download(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(client)
    await _create_chapter(client, project_id, volume_id, "Bab 1", "Isi bab 1\r\nBaris kedua", 5)
    await _create_chapter(client, project_id, volume_id, "Bab 2", "Isi bab 2", 5)

    result, job = await _run_export_job(client, session, project_id, volume_id, "pdf")

    assert result["filename"] == "Novel Uji-Lengkap-2026-07-28.pdf"
    assert result["format"] == "pdf"
    _part_path, output_path = chapter_export_service.export_file_paths(job.id, "pdf")
    assert output_path.exists()

    download_response = await client.get(
        f"/api/v1/projects/{project_id}/chapter-exports/{job.id}/download"
    )
    assert download_response.status_code == 200
    assert download_response.headers["content-type"] == "application/pdf"
    assert download_response.content[:5] == b"%PDF-"
    assert "Novel Uji-Lengkap-2026-07-28.pdf" in unquote(
        download_response.headers["content-disposition"]
    )


NON_FICTION_CHAPTER = """# Bab 1: Paradoks Kepemimpinan

### Pengantar

Paragraf dengan *miring* dan **tebal**.

| Dimensi | Nilai |
| :--- | ---: |
| **Otoritas** | 42 |

1. Butir pertama.
2. Butir kedua.

- Butir berbutir

> Kutipan penting.

---

```
[ META-LEADER ]
\u250c\u2500\u2500\u2500\u253c\u2500\u2500\u2500\u2510
```
"""


@pytest.mark.asyncio
async def test_docx_export_renders_markdown_structure_for_non_fiction(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Naskah non-fiksi harus menjadi struktur Word sungguhan, bukan tanda markdown mentah."""

    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(
        client, title="Buku Panduan", book_type="non_fiction"
    )
    await _create_chapter(
        client,
        project_id,
        volume_id,
        "Bab 1: Paradoks Kepemimpinan",
        NON_FICTION_CHAPTER,
        40,
    )

    _result, job = await _run_export_job(client, session, project_id, volume_id, "docx")
    _part_path, output_path = chapter_export_service.export_file_paths(job.id, "docx")
    document = Document(str(output_path))

    styles = {
        paragraph.text: getattr(paragraph.style, "name", None)
        for paragraph in document.paragraphs
        if paragraph.text
    }
    # Judul markdown menjadi gaya Heading, bukan paragraf berisi tanda pagar.
    assert styles["Pengantar"] == "Heading 5"
    assert styles["Butir pertama."] == "List Number"
    assert styles["Butir berbutir"] == "List Bullet"
    assert styles["Kutipan penting."] == "Quote"
    assert not any(text.startswith("#") for text in styles)
    assert not any("**" in text for text in styles)

    # Tabel markdown menjadi tabel Word sungguhan.
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Dimensi", "Nilai"]
    assert [cell.text for cell in table.rows[1].cells] == ["Otoritas", "42"]

    # Diagram bergaris kotak tetap utuh dan memakai huruf berlebar seragam.
    code_paragraph = next(p for p in document.paragraphs if "META-LEADER" in p.text)
    assert "\u250c\u2500\u2500\u2500\u253c\u2500\u2500\u2500\u2510" in code_paragraph.text
    assert {run.font.name for run in code_paragraph.runs} == {"DejaVu Sans Mono"}


@pytest.mark.asyncio
async def test_docx_export_skips_heading_that_repeats_chapter_title(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Judul bab sudah ditulis dari medannya sendiri, jadi judul markdown kembar dibuang."""

    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(
        client, title="Buku Panduan", book_type="non_fiction"
    )
    await _create_chapter(
        client,
        project_id,
        volume_id,
        "Bab 1: Paradoks Meta-Kepemimpinan - Memimpin Mereka",
        "# BAB 1: PARADOKS META-KEPEMIMPINAN: MEMIMPIN MEREKA\n\nIsi bab.",
        8,
    )

    _result, job = await _run_export_job(client, session, project_id, volume_id, "docx")
    _part_path, output_path = chapter_export_service.export_file_paths(job.id, "docx")
    document = Document(str(output_path))

    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if str(getattr(paragraph.style, "name", "")).startswith("Heading")
    ]
    # Hanya judul volume dan judul bab yang tersisa; judul markdown kembar tidak ikut tertulis.
    assert headings == ["Volume 1", "Bab 1: Paradoks Meta-Kepemimpinan - Memimpin Mereka"]


@pytest.mark.asyncio
async def test_pdf_export_accepts_box_drawing_inside_code_block(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Diagram bergaris kotak memakai font berkas, sehingga tidak lagi digagalkan penjaga aksara."""

    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(
        client, title="Buku Panduan", book_type="non_fiction"
    )
    await _create_chapter(client, project_id, volume_id, "Bab 1", NON_FICTION_CHAPTER, 40)

    result, job = await _run_export_job(client, session, project_id, volume_id, "pdf")

    assert result["format"] == "pdf"
    _part_path, output_path = chapter_export_service.export_file_paths(job.id, "pdf")
    assert output_path.read_bytes()[:5] == b"%PDF-"


@pytest.mark.asyncio
async def test_pdf_export_still_rejects_unsupported_script_in_prose(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Pengecualian font hanya berlaku di blok kode; aksara asing di prosa tetap ditolak."""

    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(
        client, title="Buku Panduan", book_type="non_fiction"
    )
    await _create_chapter(
        client, project_id, volume_id, "Bab 1", "## Subbagian\n\nParagraf \u4e2d\u6587 di prosa.", 6
    )

    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
            "format": "pdf",
        },
    )
    job = await background_service.get_job(session, created.json()["id"])
    assert job is not None
    job.status = "running"
    await session.commit()

    context = JobContext(session=session, job=job, publisher=BackgroundEventPublisher(None))
    with pytest.raises(RuntimeError):
        await dispatch_job(context)


def test_pdf_guard_accepts_latin_prose_including_typography() -> None:
    """Aksen, tanda pisah, dan kutip cerdas lazim dalam prosa Indonesia dan harus lolos."""
    accepted = "Ambigu tekad naïve café em—dash “kutip cerdas” ‘tunggal’ … 50°"
    assert pdf_renderer.find_unsupported_characters(accepted) == []


def test_pdf_guard_flags_characters_without_glyphs() -> None:
    """Aksara di luar cakupan font bawaan harus terdeteksi sebelum berkas terbentuk."""
    assert pdf_renderer.find_unsupported_characters("中文小说") == ["中", "文", "小", "说"]
    assert pdf_renderer.find_unsupported_characters("emoji 😀") == ["😀"]


@pytest.mark.asyncio
async def test_pdf_export_fails_clearly_for_unsupported_script(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Naskah non-Latin harus menggagalkan tugas PDF, bukan menghasilkan berkas berisi kotak kosong."""
    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(client)
    await _create_chapter(client, project_id, volume_id, "Bab 1", "第一章的正文", 5)
    created = await client.post(
        f"/api/v1/projects/{project_id}/chapter-exports",
        json={
            "selected_volume_ids": [volume_id],
            "included_chapter_ids": [],
            "excluded_chapter_ids": [],
            "local_date": "2026-07-28",
            "format": "pdf",
        },
    )
    job = await background_service.get_job(session, created.json()["id"])
    assert job is not None
    job.status = "running"
    await session.commit()

    context = JobContext(session=session, job=job, publisher=BackgroundEventPublisher(None))
    with pytest.raises(RuntimeError) as failure:
        await dispatch_job(context)

    assert "Word" in str(failure.value)
    part_path, output_path = chapter_export_service.export_file_paths(job.id, "pdf")
    assert not part_path.exists()
    assert not output_path.exists()


@pytest.mark.asyncio
async def test_docx_export_accepts_script_rejected_by_pdf(
    client: AsyncClient,
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Batas aksara hanya berlaku pada PDF, sedangkan DOCX menyimpan teks apa pun sebagai UTF-8."""
    async def skip_cancellation_check(_context: JobContext) -> None:
        return None

    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    monkeypatch.setattr(JobContext, "check_cancelled", skip_cancellation_check)
    project_id, volume_id = await _create_project(client)
    await _create_chapter(client, project_id, volume_id, "Bab 1", "第一章的正文", 5)

    result, job = await _run_export_job(client, session, project_id, volume_id, "docx")

    assert result["format"] == "docx"
    _part_path, output_path = chapter_export_service.export_file_paths(job.id, "docx")
    document = Document(str(output_path))
    assert "第一章的正文" in [paragraph.text for paragraph in document.paragraphs]


@pytest.mark.asyncio
async def test_cleanup_removes_expired_output_for_every_format(
    session,
    monkeypatch,
    tmp_path,
) -> None:
    """Berkas hasil setiap format wajib dikenali pembersihan berkala, bukan hanya TXT."""
    monkeypatch.setattr(chapter_export_service.settings, "chapter_exports_dir", tmp_path)
    expires_at = datetime.now(UTC) - timedelta(minutes=1)
    for index, export_format in enumerate(("txt", "docx", "pdf")):
        job = BackgroundJob(
            id=f"expired-{export_format}-{index}",
            type=chapter_export_service.EXPORT_JOB_TYPE,
            status="succeeded",
            payload_json=json.dumps({"filename": f"uji.{export_format}", "format": export_format}),
            result_json=json.dumps({"expires_at": expires_at.isoformat()}),
        )
        session.add(job)
        await session.commit()
        _part_path, output_path = chapter_export_service.export_file_paths(job.id, export_format)
        output_path.write_text("kedaluwarsa", encoding="utf-8")

        assert await chapter_export_service.cleanup_chapter_export_files(session) == 1
        assert not output_path.exists()
